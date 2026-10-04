"""Resume text extraction (PDF / DOCX / TXT) and structured field parsing.
Fields that cannot be found are returned as None / empty lists - nothing is invented."""
import io
import re
from collections import Counter
from .education import education_level, extract_degrees, LEVEL_NAMES
from .skill_extractor import extract_skills, ALL_SKILLS
from .text_processor import tokenize, remove_stopwords

MAX_BYTES = 10 * 1024 * 1024


class ResumeParseError(Exception):
    """User-friendly error for unreadable resume files."""


def extract_text(data: bytes, filename: str) -> str:
    if not data:
        raise ResumeParseError("The uploaded file is empty.")
    if len(data) > MAX_BYTES:
        raise ResumeParseError("File is larger than 10 MB.")
    ext = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    try:
        if ext == "pdf":
            import pymupdf
            with pymupdf.open(stream=data, filetype="pdf") as doc:
                text = "\n".join(page.get_text() for page in doc)
        elif ext == "docx":
            import docx
            d = docx.Document(io.BytesIO(data))
            parts = [p.text for p in d.paragraphs]
            for t in d.tables:
                for row in t.rows:
                    parts.append(" | ".join(c.text for c in row.cells))
            text = "\n".join(parts)
        elif ext == "txt":
            for enc in ("utf-8", "latin-1"):
                try:
                    text = data.decode(enc)
                    break
                except UnicodeDecodeError:
                    continue
        else:
            raise ResumeParseError(f"Unsupported file type '.{ext}'. Please upload PDF, DOCX or TXT.")
    except ResumeParseError:
        raise
    except Exception as e:
        raise ResumeParseError(f"Could not read this {ext.upper()} file (it may be corrupted or password-protected): {e}")
    text = text.replace("\x00", " ").strip()
    if len(text) < 30:
        raise ResumeParseError("No readable text found in the resume (scanned/image-only PDFs are not supported - no OCR).")
    return text


_SECTION_HEADERS = {
    "education": r"education|academic|qualification",
    "experience": r"(?:professional |work |employment )?experience|employment|work history|internship",
    "skills": r"(?:technical |key |core )?skills|technologies|competenc",
    "projects": r"projects?",
    "certifications": r"certifications?|courses",
    "summary": r"summary|objective|profile|about me",
}
_MONTHS = "jan|feb|mar|apr|may|jun|jul|aug|sep|sept|oct|nov|dec"
_DATE_RANGE = re.compile(
    rf"((?:(?:{_MONTHS})[a-z]*\.?\s+)?(?:19|20)\d{{2}})\s*(?:-|–|—|to)\s*((?:(?:{_MONTHS})[a-z]*\.?\s+)?(?:(?:19|20)\d{{2}})|present|current|now|ongoing)", re.I)


def split_sections(text: str) -> dict:
    sections, current, buf = {}, "header", []
    for line in text.splitlines():
        s = line.strip().strip(":").strip()
        hit = None
        if 0 < len(s) <= 40 and not re.search(r"[|@]", s):
            for name, pat in _SECTION_HEADERS.items():
                if re.fullmatch(rf"(?:{pat})(?:\s*(?:&|and)\s*\w+)?", s, re.I):
                    hit = name
                    break
        if hit:
            sections[current] = "\n".join(buf)
            current, buf = hit, []
        else:
            buf.append(line)
    sections[current] = "\n".join(buf)
    return {k: v.strip() for k, v in sections.items() if v.strip()}


def _parse_date(s: str, end=False):
    from datetime import date
    s = s.strip().lower()
    if s in ("present", "current", "now", "ongoing"):
        t = date.today()
        return t.year + (t.month - 1) / 12
    y = int(re.search(r"(?:19|20)\d{2}", s).group())
    m = re.match(rf"({_MONTHS})", s)
    mon = ["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"].index(m.group(1)) if m else (11 if end else 0)
    return y + mon / 12


def estimate_experience_years(text: str, sections: dict):
    """Returns (years or None, method). Explicit statement wins; else merged date ranges in the experience section."""
    m = re.findall(r"(\d+(?:\.\d+)?)\s*\+?\s*(?:years?|yrs?)\s*(?:of\s*)?(?:professional\s*|relevant\s*|total\s*)?experience", text, re.I)
    if m:
        return max(float(x) for x in m), "stated in resume"
    exp = sections.get("experience", "")
    ranges = []
    for a, b in _DATE_RANGE.findall(exp):
        try:
            s, e = _parse_date(a), _parse_date(b, end=True)
            if 0 <= e - s < 40:
                ranges.append((s, e))
        except Exception:
            continue
    if not ranges:
        return None, None
    ranges.sort()
    merged = [list(ranges[0])]
    for s, e in ranges[1:]:
        if s <= merged[-1][1]:
            merged[-1][1] = max(merged[-1][1], e)
        else:
            merged.append([s, e])
    return round(sum(e - s for s, e in merged), 1), "estimated from date ranges"


def _guess_name(text: str):
    for line in text.splitlines()[:8]:
        s = line.strip()
        if not s or re.search(r"@|\d{5,}|http|www\.|resume|curriculum|vitae", s, re.I):
            continue
        words = s.split()
        if 1 < len(words) <= 4 and all(re.fullmatch(r"[A-Za-z.\-']+", w) for w in words) and \
                not re.fullmatch("|".join(_SECTION_HEADERS.values()), s, re.I):
            return s.title() if s.isupper() else s
    return None


def extract_keywords(text: str, top: int = 15) -> list:
    """Frequent non-skill content words (supporting keywords)."""
    skill_words = {w for s in ALL_SKILLS for w in s.lower().split()}
    toks = [t for t in remove_stopwords(tokenize(text)) if len(t) > 3 and t not in skill_words and not re.search(r"\d", t) and "." not in t]
    return [w for w, _ in Counter(toks).most_common(top)]


def parse_resume(text: str) -> dict:
    """Parse resume text into a structured dict. Missing information stays None / []."""
    sections = split_sections(text)
    email = re.search(r"[\w.+-]+@[\w-]+\.[\w.-]+", text)
    phone = re.search(r"(?:\+?\d{1,3}[\s-]?)?(?:\(?\d{3,5}\)?[\s-]?)?\d{3,5}[\s-]?\d{4,5}", text)
    phone_s = phone.group().strip() if phone and len(re.sub(r"\D", "", phone.group())) >= 10 else None
    edu_text = sections.get("education", "")
    degrees = extract_degrees(edu_text or text)
    level = education_level(edu_text) if edu_text else education_level(" ".join(degrees))
    years, method = estimate_experience_years(text, sections)
    name = _guess_name(text)
    match_text = text                       # used for similarity: identity/contact details are excluded
    for pii in filter(None, [name, email.group() if email else None, phone_s]):
        match_text = match_text.replace(pii, " ")
    return {
        "name": name, "match_text": match_text,
        "email": email.group() if email else None,
        "phone": phone_s,
        "linkedin": (re.search(r"linkedin\.com/[\w\-/]+", text, re.I) or [None])[0] if re.search(r"linkedin\.com", text, re.I) else None,
        "github": (re.search(r"github\.com/[\w\-]+", text, re.I) or [None])[0] if re.search(r"github\.com", text, re.I) else None,
        "skills": extract_skills(text),
        "education": {"degrees": degrees, "level": level, "level_name": LEVEL_NAMES.get(level) if level is not None else None,
                      "section_text": edu_text[:600] or None},
        "experience": {"years": years, "method": method, "section_text": sections.get("experience", "")[:800] or None},
        "keywords": extract_keywords(text),
        "sections_found": [k for k in sections if k != "header"],
        "n_words": len(text.split()),
        "raw_text": text,
    }
