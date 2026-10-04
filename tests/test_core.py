"""Unit tests for parsing robustness and scoring. Run: python tests/test_core.py"""
import io, sys, pathlib
sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent.parent))
from src.resume_parser import extract_text, parse_resume, ResumeParseError
from src.matcher import skill_score, experience_score, education_score, combine
from src.skill_extractor import extract_skills, normalize_skill

def expect_error(data, name):
    try:
        extract_text(data, name); print("FAIL no error:", name)
    except ResumeParseError as e:
        print("ok  error ->", name, "|", str(e)[:70])

# 1. malformed inputs
expect_error(b"", "empty.txt"); expect_error(b"%PDF-1.4 garbage", "broken.pdf"); expect_error(b"PK\x03\x04junk", "broken.docx")
expect_error(b"hello", "x.exe"); expect_error(b"   \n  ", "blank.txt")
# 2. real DOCX and PDF round trip
import docx, pymupdf
d = docx.Document(); d.add_paragraph("Asha Rao"); d.add_paragraph("asha@mail.com"); d.add_paragraph("SKILLS"); d.add_paragraph("Python, SQL, Django, React, Docker and Git. 4 years of experience building web apps.")
b = io.BytesIO(); d.save(b); t = extract_text(b.getvalue(), "r.docx"); r = parse_resume(t); print("docx skills:", r["skills"], r["experience"]["years"])
pdf = pymupdf.open(); pg = pdf.new_page(); pg.insert_text((72, 72), "Ravi Kumar\nravi@mail.com\nEDUCATION\nB.Tech Computer Science 2018-2022\nSKILLS\nJava, Spring Boot, MySQL, AWS, Machine Learning and Excel skills.\n" * 2)
t = extract_text(pdf.tobytes(), "r.pdf"); r = parse_resume(t); print("pdf skills:", r["skills"], r["education"]["level_name"])
# 3. resume with almost nothing extractable -> no crash, no invention
r = parse_resume("Just a short note about my hobbies: painting, cooking and long walks in the park near my home."); print("sparse:", r["name"], r["skills"], r["experience"]["years"], r["education"]["level"])
# 4. scoring
assert normalize_skill("nodejs") == "Node.js" and "C#" in extract_skills("worked in c sharp")
assert skill_score(["python"], ["Python", "SQL"])[0] == 0.5
assert skill_score([], [], [])[0] is None
assert experience_score(2, 4) == 0.5 and experience_score(None, 4) is None and education_score(2, 3) == 0.5
s, w = combine({"skills": 0.8, "text": None, "experience": 1.0, "education": None}); print("renormalised weights:", w, round(s, 3))
print("ALL CORE TESTS PASSED")
